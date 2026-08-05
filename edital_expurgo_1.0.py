# importações de bibliotecas
import locale
import os
import re
from datetime import datetime

import pandas as pd
from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from num2words import num2words
from openpyxl import load_workbook

# constantes de caminho e medidas
# pasta do SharePoint (DETRAN - Divisão de Gestão Documental) sincronizada localmente
PASTA = r"C:\Users\julia.santos\PRODESP\DETRAN - DIVISÃO DE GESTÃO DOCUMENTAL - Documentos\Editais de Eliminação de Documentos"
ARQUIVO = os.path.join(PASTA, "Relacao de Expurgo.xlsx")
EDITAIS = os.path.join(PASTA, "Editais Elaborados")
MODELO = os.path.join(PASTA, "Modelos")

MODELO_CABECALHO = os.path.join(MODELO, "modelo_cabeçalho.txt")
MODELO_DETALHAMENTO = os.path.join(MODELO, "modelo_detalhamento.txt")
MODELO_RODAPE = os.path.join(MODELO, "modelo_rodapé.txt")

METRAGEM_MEDIDA = 0.14  # metros lineares

# funções auxiliares
def limpar_nome(nome):
    return re.sub(r'[\\/*?:"<>|]', "_", str(nome))


def capitalizar_personalizado(texto):
    conectores = {
        'de', 'do', 'da', 'dos', 'das', 'e',
        'em', 'com', 'no', 'na', 'nos', 'nas',
        'a', 'o', 'as', 'os'
    }

    palavras = texto.lower().split()
    resultado = []

    for i, palavra in enumerate(palavras):
        if i == 0 or palavra not in conectores:
            resultado.append(palavra.capitalize())
        else:
            resultado.append(palavra)

    return ' '.join(resultado)


def extrair_chave_ordenacao(codigo):
    return [int(p) for p in str(codigo).split('.') if p.isdigit()]


# código principal
dataframe = pd.read_excel(
    ARQUIVO,
    sheet_name="Edital de Caixa",
    engine='openpyxl'
)

limpar_colunas = [
    'Função',
    'Subfunção',
    'Atividade',
    'Série documental',
    'Observações complementares'
]

for campo in limpar_colunas + ['Região Administrativa', 'Município']:
    dataframe[campo] = dataframe[campo].astype(str).str.strip()

# filtra apenas as linhas com "Criar edital"
df_criar_edital = dataframe[
    dataframe['Status Edital'].str.contains(
        "Criar edital",
        case=False,
        na=False
    )
]

if df_criar_edital.empty:
    print("Nenhum edital a ser criado.")
    exit()

solicitacao = df_criar_edital.groupby(
    ['N° Processo SEI', 'Região Administrativa', 'Município']
)

for (processo, regiao, municipio), grupo in solicitacao:

    # cabeçalho
    with open(MODELO_CABECALHO, "r", encoding="utf-8") as modelo:

        locale.setlocale(locale.LC_ALL, 'pt_BR.UTF-8')

        data_edital = datetime.now().strftime(
            "%d de %B de %Y"
        ).upper()

        regiao = str(
            grupo['Região Administrativa'].iloc[0]
        ).strip()

        municipio = str(
            grupo['Município'].iloc[0]
        ).strip()

        modelo_cabecalho = modelo.read()

        substituicoes_cabecalho = {
            "{data_edital}": data_edital,
            "{regiao}": regiao,
            "{municipio}": municipio
        }

        cabecalho = modelo_cabecalho

        for chave, valor in substituicoes_cabecalho.items():
            cabecalho = cabecalho.replace(chave, valor)

    # limpeza das colunas
    for campo in limpar_colunas:
        grupo.loc[:, campo] = (
            grupo[campo]
            .astype(str)
            .str.strip()
        )

    # ordenação
    grupo.loc[:, 'chave_ordenacao'] = (
        grupo['Série documental']
        .apply(extrair_chave_ordenacao)
    )

    # agrupamento
    detalhamento = grupo.groupby([
        'Função',
        'Subfunção',
        'Atividade',
        'Série documental',
        'Descrição documental',
        'Data Limite',
        'Quantidade',
        'Observações complementares'
    ]).first().reset_index()

    detalhamento['chave_ordenacao'] = (
        detalhamento['Série documental']
        .apply(extrair_chave_ordenacao)
    )

    detalhamento = detalhamento.sort_values(
        by='chave_ordenacao'
    )

    # detalhamento
    itens_detalhamento = []

    for _, row in detalhamento.iterrows():

        with open(
            MODELO_DETALHAMENTO,
            "r",
            encoding="utf-8"
        ) as detalhamento_modelo:

            detalhamento_item = detalhamento_modelo.read()

            qtde_caixas = int(row['Quantidade'])

            caixas_extenso = num2words(
                qtde_caixas,
                lang='pt_BR',
                gender='f'
            )

            if qtde_caixas == 1:
                caixas_extenso = f"({caixas_extenso}) Caixa"
            else:
                caixas_extenso = f"({caixas_extenso}) Caixas"

            substituicoes_detalhamento = {
                "{funcao}": capitalizar_personalizado(
                    row['Função']
                ).replace('_x000D_', ''),

                "{subfuncao}": row['Subfunção']
                .replace('_x000D_', ' ')
                .capitalize(),

                "{atividade}": row['Atividade']
                .replace('_x000D_', ' '),

                "{serie_documental}": row['Série documental']
                .replace('_x000D_', ' '),

                "{descricao_documental}": row['Descrição documental']
                .replace('_x000D_', ''),

                "{data_limite}": str(row['Data Limite'])
                .replace('_x000D_', ' '),

                "{qtde_caixas}": str(qtde_caixas).zfill(2),

                "{caixas_extenso}": caixas_extenso,

                "{observacoes_complementares}": (
                    row['Observações complementares']
                    .replace('_x000D_', ' ')
                    .replace('nan', '')
                )
            }

            for chave, valor in substituicoes_detalhamento.items():
                detalhamento_item = detalhamento_item.replace(
                    chave,
                    valor
                )

            itens_detalhamento.append(detalhamento_item)

    DETALHAMENTO_COMPLETO = "\n".join(itens_detalhamento)

    # rodapé
    total_caixas = int(grupo['Quantidade'].sum())

    total_caixas_extenso = num2words(
        total_caixas,
        lang='pt_BR',
        gender='f'
    )

    total_metros_lineares = (
        total_caixas * METRAGEM_MEDIDA
    )

    if total_caixas == 1:
        total_caixas_extenso = (
            f"({total_caixas_extenso}) Caixa"
        )
    else:
        total_caixas_extenso = (
            f"({total_caixas_extenso}) Caixas"
        )

    with open(MODELO_RODAPE, "r", encoding="utf-8") as rodape:

        rodape_texto = rodape.read()

        substituicoes_rodape = {
            "{total_caixas}": str(total_caixas),

            "{total_caixas_extenso}": total_caixas_extenso,

            "{total_metros_lineares}":
                f"{total_metros_lineares:.2f}"
        }

    for chave, valor in substituicoes_rodape.items():
        rodape_texto = rodape_texto.replace(
            chave,
            valor
        )

    edital_rodape = rodape_texto

    # conteúdo final
    SALVAR_EDITAL = f"""{cabecalho}
{DETALHAMENTO_COMPLETO}
{edital_rodape}"""

    # salva documento
    nome_arquivo = (
        f"Edital_{limpar_nome(municipio)}_"
        f"{limpar_nome(processo)}.docx"
    )

    caminho_arquivo = os.path.join(
        EDITAIS,
        nome_arquivo
    )

    documento = Document()

    for linha in SALVAR_EDITAL.split('\n'):
        paragrafo = documento.add_paragraph(linha)
        paragrafo.alignment = (
            WD_PARAGRAPH_ALIGNMENT.JUSTIFY
        )

    documento.save(caminho_arquivo)

    # atualiza excel
    wb = load_workbook(ARQUIVO)

    aba = wb["Edital de Caixa"]

    for idx in grupo.index:
        excel_row = idx + 2
        aba[f"O{excel_row}"] = "Edital criado"

    wb.save(ARQUIVO)

print("Editais criados com sucesso!")